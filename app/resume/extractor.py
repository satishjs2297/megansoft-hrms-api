import os
import tempfile
from typing import Tuple
import pdfplumber
from docx import Document
from app.config import get_settings

settings = get_settings()

def extract_text_from_pdf(file_bytes: bytes, filename: str) -> str:
    full_text = []
    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
        tmp.write(file_bytes)
        tmp_path = tmp.name
    try:
        with pdfplumber.open(tmp_path) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text and text.strip():
                    full_text.append(text)
        if not full_text and settings.use_ocr:
            full_text = [_ocr_pdf(tmp_path)]
    finally:
        os.unlink(tmp_path)
    return "\n\n".join(full_text)

def extract_text_from_docx(file_bytes: bytes) -> str:
    full_text = []
    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
        tmp.write(file_bytes)
        tmp_path = tmp.name
    try:
        doc = Document(tmp_path)
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                full_text.append(paragraph.text)
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    full_text.append(" | ".join(row_text))
    finally:
        os.unlink(tmp_path)
    return "\n".join(full_text)

def _ocr_pdf(file_path: str) -> str:
    try:
        import easyocr
        from pdf2image import convert_from_path
        reader = easyocr.Reader(['en'])
        images = convert_from_path(file_path)
        texts = []
        for image in images:
            result = reader.readtext(image, detail=0)
            texts.append("\n".join(result))
        return "\n\n".join(texts)
    except Exception as e:
        return ""

def extract_text(file_bytes: bytes, filename: str) -> Tuple[str, str]:
    """Returns (extracted_text, file_type)"""
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes, filename), "pdf"
    elif lower.endswith((".docx", ".doc")):
        return extract_text_from_docx(file_bytes), "docx"
    elif lower.endswith(".txt"):
        return file_bytes.decode("utf-8", errors="ignore"), "txt"
    else:
        raise ValueError(f"Unsupported file type: {filename}")
